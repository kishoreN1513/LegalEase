from io import BytesIO
from docx import Document
from docx.shared import Pt
from fpdf import FPDF

def make_txt(text):
    return text.encode("utf-8")

def make_docx(text, title="LegalEase Document"):
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()

def make_pdf(text, title="LegalEase Document"):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, title, ln=True, align="C")
    pdf.ln(5)
    pdf.set_font("Helvetica", size=10)
    safe = text.encode("latin-1", "replace").decode("latin-1")
    for line in safe.split("\n"):
        if line.strip():
            pdf.multi_cell(0, 6, line)
        else:
            pdf.ln(3)
    return bytes(pdf.output())
